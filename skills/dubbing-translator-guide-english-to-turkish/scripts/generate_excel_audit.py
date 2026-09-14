#!/usr/bin/env python3
"""
Dubbing Excel Workbook Generator Script
Used by dubbing-translator-guide-english-to-turkish skill.

Generates formatted 5-column comparative Excel audit sheets (.xlsx).
Columns: English Dialogue | N_en | Phonetic Turkish Translation | N_tr | Absolute Difference (|N_tr - N_en|)
"""

import sys
import os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
import syllables

def count_syllables_en(text: str) -> int:
    words = text.split()
    return sum(syllables.estimate(w) for w in words)

def count_syllables_tr(text: str) -> int:
    vowels = "aeıioöuüAEIİOÖUÜ"
    return sum(1 for char in text if char in vowels)

def create_excel_workbook(en_file_path: str, tr_file_path: str, output_excel_path: str):
    if not os.path.exists(en_file_path):
        print(f"Error: English source file not found at '{en_file_path}'")
        sys.exit(1)
    if not os.path.exists(tr_file_path):
        print(f"Error: Turkish dubbing file not found at '{tr_file_path}'")
        sys.exit(1)

    with open(en_file_path, "r", encoding="utf-8") as f:
        en_lines = [l.strip() for l in f if l.strip()]

    with open(tr_file_path, "r", encoding="utf-8") as f:
        tr_lines = [l.strip() for l in f if l.strip()]

    if len(en_lines) != len(tr_lines):
        print(f"Warning: Line count mismatch! EN={len(en_lines)} vs TR={len(tr_lines)}")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dubbing Lip-Sync Audit"

    headers = [
        "English Dialogue",
        "English Syllables (N_en)",
        "Phonetic Turkish Dubbing Translation",
        "Turkish Syllables (N_tr)",
        "Absolute Difference (|N_tr - N_en|)"
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for en, tr in zip(en_lines, tr_lines):
        n_en = count_syllables_en(en)
        n_tr = count_syllables_tr(tr)
        diff = abs(n_tr - n_en)
        ws.append([en, n_en, tr, n_tr, diff])

    for row in ws.iter_rows(min_row=2, max_row=len(en_lines)+1, min_col=1, max_col=5):
        row[0].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        row[1].alignment = Alignment(horizontal="center", vertical="center")
        row[2].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        row[3].alignment = Alignment(horizontal="center", vertical="center")
        row[4].alignment = Alignment(horizontal="center", vertical="center")

    ws.column_dimensions["A"].width = 50
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 55
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 18

    wb.save(output_excel_path)
    print(f"[SUCCESS] Excel audit workbook saved to '{output_excel_path}'")

if __name__ == "__main__":
    if len(sys.argv) < 4:
        print("Usage: python3 generate_excel_audit.py <english_source.md> <turkish_dubbing.md> <output_audit.xlsx>")
        sys.exit(1)
    create_excel_workbook(sys.argv[1], sys.argv[2], sys.argv[3])
